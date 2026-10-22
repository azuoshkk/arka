from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".rst",
    ".pdf",
    ".docx",

    # código
    ".py",
    ".h",
    ".hpp",
    ".c",
    ".cc",
    ".cpp",
    ".ts",
    ".tsx",
    ".js",
    ".jsx",

    # configuração
    ".yaml",
    ".yml",
    ".json",
    ".xml",
    ".csv",
    ".toml",
    ".ini",
    ".cfg",
    ".conf",
    ".msg",

    # shell
    ".sh",
    ".bash",
}


SPECIAL_FILENAMES = {
    "dockerfile",
    "makefile",
}


def is_supported_file(path: Path) -> bool:

    name = path.name.lower()

    if path.suffix.lower() in SUPPORTED_EXTENSIONS:
        return True

    if name in SPECIAL_FILENAMES:
        return True

    # A curadoria pode gerar Dockerfile__2
    if name.startswith("dockerfile__"):
        return True

    return False


def _load_pdf(path: Path):

    reader = PdfReader(str(path))

    documents = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):

        text = (
            page.extract_text()
            or ""
        ).strip()

        if not text:
            continue

        documents.append({
            "text": text,
            "page": page_number,
        })

    return documents


def _load_docx(path: Path):

    document = Document(str(path))

    parts = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                parts.append(
                    " | ".join(cells)
                )

    text = "\n".join(parts).strip()

    if not text:
        return []

    return [{
        "text": text,
        "page": None,
    }]


def _load_plain_text(path: Path):

    text = path.read_text(
        encoding="utf-8",
        errors="ignore",
    ).strip()

    if not text:
        return []

    return [{
        "text": text,
        "page": None,
    }]


def load_file(path: Path):

    if not is_supported_file(path):
        return []

    extension = path.suffix.lower()

    try:

        if extension == ".pdf":
            return _load_pdf(path)

        if extension == ".docx":
            return _load_docx(path)

        return _load_plain_text(path)

    except Exception as exc:

        print(
            f"⚠️ Erro lendo {path}: {exc}"
        )

        return []
